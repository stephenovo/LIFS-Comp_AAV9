"""Audits for the public Fit4Function release.

The public release mixes sequence-linked CSV files with anonymized Excel
workbooks.  This module makes that distinction explicit because a table can be
useful for reproducing a figure without being sufficient for training a new
sequence-to-function model.
"""

from __future__ import annotations

import hashlib
import itertools
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.mixture import GaussianMixture

from aav9_sma.data.audit import is_valid_peptide

SEQUENCE_FILES = (
    "modeling_library_production_fitness.csv",
    "assessment_library_production_fitness.csv",
    "fit4function_library_screens.csv",
    "nnk_library_top_production_fitness_240k.csv",
)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _numeric_summary(frame: pd.DataFrame) -> dict[str, dict[str, float | int | None]]:
    summary: dict[str, dict[str, float | int | None]] = {}
    for column in frame.select_dtypes(include="number"):
        values = frame[column]
        finite = values.replace([np.inf, -np.inf], np.nan)
        summary[column] = {
            "finite_count": int(finite.notna().sum()),
            "nan_count": int(values.isna().sum()),
            "positive_inf_count": int(np.isposinf(values).sum()),
            "negative_inf_count": int(np.isneginf(values).sum()),
            "min": float(finite.min()) if finite.notna().any() else None,
            "median": float(finite.median()) if finite.notna().any() else None,
            "max": float(finite.max()) if finite.notna().any() else None,
        }
    return summary


def _csv_summary(path: Path) -> dict[str, object]:
    frame = pd.read_csv(path)
    result: dict[str, object] = {
        "bytes": path.stat().st_size,
        "sha256": _sha256(path),
        "rows": len(frame),
        "columns": list(frame.columns),
        "numeric": _numeric_summary(frame),
        "has_amino_acid_sequence": "AA" in frame,
    }
    if "AA" in frame:
        amino_acids = frame["AA"]
        result.update(
            {
                "unique_amino_acid_sequences": int(amino_acids.nunique(dropna=True)),
                "duplicate_amino_acid_rows": int(amino_acids.duplicated().sum()),
                "valid_7mer_rows": int(amino_acids.map(is_valid_peptide).sum()),
            }
        )
    return result


def _workbook_structure(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [
            {
                "name": sheet.title,
                "state": sheet.sheet_state,
                "rows_including_header": sheet.max_row,
                "columns": sheet.max_column,
            }
            for sheet in workbook.worksheets
        ]
        defined_names = list(workbook.defined_names)
    finally:
        workbook.close()
    return {"sheets": sheets, "defined_names": defined_names}


def _production_mixture_summary(path: Path) -> dict[str, object]:
    frame = pd.read_csv(path, usecols=["Label", "Production"])
    raw = frame.loc[frame["Label"].eq("Designed"), "Production"].to_numpy(dtype=float)
    values = np.log2(raw[np.isfinite(raw) & (raw > 0)]).reshape(-1, 1)
    mixture = GaussianMixture(n_components=2, n_init=20, random_state=42).fit(values)
    order = np.argsort(mixture.means_.ravel())
    grid = np.linspace(float(values.min()), float(values.max()), 100_000).reshape(-1, 1)
    fit_probability = mixture.predict_proba(grid)[:, order[1]]
    boundary_index = int(np.argmin(np.abs(fit_probability - 0.5)))
    boundary = float(grid[boundary_index, 0])
    return {
        "method": "two-component-Gaussian-mixture-on-finite-designed-log2-production",
        "rows": len(values),
        "non_fit_component": {
            "mean": float(mixture.means_[order[0], 0]),
            "sd": float(np.sqrt(mixture.covariances_[order[0], 0, 0])),
            "weight": float(mixture.weights_[order[0]]),
        },
        "fit_component": {
            "mean": float(mixture.means_[order[1], 0]),
            "sd": float(np.sqrt(mixture.covariances_[order[1], 0, 0])),
            "weight": float(mixture.weights_[order[1]]),
        },
        "posterior_equal_boundary_log2": boundary,
        "posterior_equal_boundary_raw_ratio": float(2**boundary),
        "fraction_of_finite_rows_above_boundary": float((values[:, 0] >= boundary).mean()),
        "caveat": "An audit estimate, not a pre-approved project gate.",
    }


def _invivo_measurement_summary(path: Path) -> dict[str, object]:
    sheets = pd.read_excel(path, sheet_name=None)
    organs: dict[str, object] = {}
    row_aligned_means: dict[str, pd.Series] = {}
    for organ, frame in sheets.items():
        values = frame.iloc[:, 1:].apply(pd.to_numeric, errors="coerce")
        correlations = values.corr(min_periods=100)
        upper = correlations.where(np.triu(np.ones(correlations.shape), 1).astype(bool)).stack()
        suffix = frame["SequenceID"].astype(str).str.extract(r"_(\d+)$", expand=False)
        sequential_ids = bool(
            suffix.notna().all()
            and np.array_equal(suffix.astype(int).to_numpy(), np.arange(1, len(frame) + 1))
        )
        organs[organ] = {
            "rows": len(frame),
            "columns": list(frame.columns),
            "sequence_id_example": str(frame["SequenceID"].iloc[0]),
            "contains_amino_acid_sequence": "AA" in frame,
            "sequential_anonymous_ids": sequential_ids,
            "non_null_by_animal": {
                column: int(values[column].notna().sum()) for column in values
            },
            "pairwise_animal_pearson_median": float(upper.median()),
            "pairwise_animal_pearson_min": float(upper.min()),
            "pairwise_animal_pearson_max": float(upper.max()),
        }
        row_aligned_means[organ] = values.mean(axis=1, skipna=True)

    aligned = pd.DataFrame(row_aligned_means)
    aligned_correlations = aligned.corr(min_periods=1000)
    return {
        "bytes": path.stat().st_size,
        "sha256": _sha256(path),
        "structure": _workbook_structure(path),
        "organs": organs,
        "row_aligned_cross_organ_correlations": {
            row: {column: float(value) for column, value in values.items()}
            for row, values in aligned_correlations.to_dict(orient="index").items()
        },
        "row_aligned_complete_all_organs": int(aligned.notna().all(axis=1).sum()),
        "sequence_linkage_status": "anonymous-organ-specific-ids-no-AA-column",
    }


def _invivo_prediction_summary(path: Path) -> dict[str, object]:
    sheets = pd.read_excel(path, sheet_name=None)
    organs: dict[str, object] = {}
    for organ, frame in sheets.items():
        paired = frame[["Measured", "Predicted"]].replace([np.inf, -np.inf], np.nan).dropna()
        organs[organ] = {
            "rows": len(frame),
            "sequence_id_example": str(frame["SequenceID"].iloc[0]),
            "contains_amino_acid_sequence": "AA" in frame,
            "paired_measurements": len(paired),
            "measured_vs_predicted_pearson": float(paired.corr().iloc[0, 1]),
        }
    return {
        "bytes": path.stat().st_size,
        "sha256": _sha256(path),
        "structure": _workbook_structure(path),
        "organs": organs,
        "sequence_linkage_status": "anonymous-organ-specific-ids-no-AA-column",
    }


def audit_official_release(repository_root: str | Path) -> dict[str, object]:
    """Return a reproducible audit of a checked-out Fit4Function repository."""
    root = Path(repository_root)
    data_root = root / "data"
    missing = [name for name in SEQUENCE_FILES if not (data_root / name).exists()]
    if missing:
        raise FileNotFoundError(f"Missing Fit4Function files: {missing}")

    csv_files = {name: _csv_summary(data_root / name) for name in SEQUENCE_FILES}
    sequence_sets = {
        name: set(pd.read_csv(data_root / name, usecols=["AA"])["AA"].dropna())
        for name in SEQUENCE_FILES
    }
    intersections = {
        f"{left}__{right}": len(sequence_sets[left] & sequence_sets[right])
        for left, right in itertools.combinations(SEQUENCE_FILES, 2)
    }

    return {
        "release_root": str(root),
        "csv_files": csv_files,
        "sequence_intersections": intersections,
        "production_mixture": _production_mixture_summary(
            data_root / "modeling_library_production_fitness.csv"
        ),
        "invivo_measurements": _invivo_measurement_summary(
            data_root / "fit4function_library_invivo.xlsx"
        ),
        "invivo_predictions": _invivo_prediction_summary(
            data_root / "fit4function_library_invivo_predictions.xlsx"
        ),
        "training_readiness": {
            "production_sequence_model": "ready",
            "liver_and_human_liver_sequence_models": "ready-from-100k-sample",
            "brain_sequence_model": "blocked-in-processed-release",
            "spinal_cord_sequence_model": "blocked-in-processed-release",
            "heart_sequence_model": "blocked-in-processed-release",
            "kidney_sequence_model": "blocked-in-processed-release",
            "reason": (
                "Processed multi-organ workbooks expose only organ-specific anonymous IDs; "
                "they contain neither AA sequences nor a shared lookup table."
            ),
        },
    }
